"""Text and metadata extraction for PDF, DOCX, TXT, PPTX and XLSX files.

Each extractor returns an ``ExtractedDocument`` with one entry per page,
slide, sheet or section. When text cannot be extracted reliably the document
is marked ``needs_visual_review`` rather than discarded.
"""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import datetime

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".pptx", ".xlsx"}

# Below this many alphabetic characters per page a PDF page is treated as
# image-only or unreliable.
MIN_ALPHA_PER_PAGE = 40
MIN_ALPHA_TOTAL = 80


@dataclass
class ExtractedDocument:
    pages: list[str] = field(default_factory=list)
    page_label: str = "page"  # page, slide, sheet or section
    paged: bool = True  # False when page numbers are not meaningful
    metadata_author: str = ""
    metadata_title: str = ""
    metadata_date: str = ""
    status: str = "ok"  # ok or needs_visual_review
    status_detail: str = ""

    @property
    def text(self) -> str:
        return "\n\n".join(self.pages)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text.lower())).strip()


def _alpha_count(text: str) -> int:
    return sum(ch.isalpha() for ch in text)


def _pdf_date(value: str) -> str:
    match = re.match(r"D:(\d{4})(\d{2})?(\d{2})?", value or "")
    if not match:
        return ""
    year, month, day = match.group(1), match.group(2) or "01", match.group(3) or "01"
    return f"{year}-{month}-{day}"


def _iso(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return ""


def extract_pdf(data: bytes) -> ExtractedDocument:
    import pymupdf

    doc = ExtractedDocument(page_label="page")
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        meta = pdf.metadata or {}
        doc.metadata_author = (meta.get("author") or "").strip()
        doc.metadata_title = (meta.get("title") or "").strip()
        doc.metadata_date = _pdf_date(meta.get("creationDate") or meta.get("modDate") or "")
        weak_pages = 0
        for page in pdf:
            text = page.get_text("text") or ""
            doc.pages.append(text)
            if _alpha_count(text) < MIN_ALPHA_PER_PAGE:
                weak_pages += 1
        total = len(doc.pages)
    if total == 0:
        doc.status, doc.status_detail = "needs_visual_review", "The PDF has no pages."
    elif _alpha_count(doc.text) < MIN_ALPHA_TOTAL:
        doc.status = "needs_visual_review"
        doc.status_detail = "No usable text layer. The PDF may be scanned or image-based."
    elif weak_pages / total > 0.5:
        doc.status = "needs_visual_review"
        doc.status_detail = f"{weak_pages} of {total} pages have little or no extractable text."
    return doc


def extract_docx(data: bytes) -> ExtractedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    props = document.core_properties
    doc = ExtractedDocument(page_label="section", paged=False)
    doc.metadata_author = (props.author or "").strip()
    doc.metadata_title = (props.title or "").strip()
    doc.metadata_date = _iso(props.created) or _iso(props.modified)

    blocks: list[str] = []
    current: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") and current:
            blocks.append("\n".join(current))
            current = []
        if text:
            prefix = "- " if "list" in style else ""
            current.append(prefix + text)
        elif current and current[-1] != "":
            current.append("")
    if current:
        blocks.append("\n".join(current))
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            deduped = [c for i, c in enumerate(cells) if c and (i == 0 or c != cells[i - 1])]
            if deduped:
                rows.append(" | ".join(deduped))
        if rows:
            blocks.append("\n".join(rows))
    doc.pages = [b.strip() for b in blocks if b.strip()] or [""]
    if _alpha_count(doc.text) < MIN_ALPHA_TOTAL:
        doc.status = "needs_visual_review"
        doc.status_detail = "The Word document contains little or no text. Content may be in images."
    return doc


def extract_pptx(data: bytes) -> ExtractedDocument:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    props = prs.core_properties
    doc = ExtractedDocument(page_label="slide")
    doc.metadata_author = (props.author or "").strip()
    doc.metadata_title = (props.title or "").strip()
    doc.metadata_date = _iso(props.created) or _iso(props.modified)
    for slide in prs.slides:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Notes: " + notes)
        doc.pages.append("\n".join(parts))
    if not doc.pages or _alpha_count(doc.text) < MIN_ALPHA_TOTAL:
        doc.status = "needs_visual_review"
        doc.status_detail = "The slides contain little or no text. Content may be in images."
    return doc


def extract_xlsx(data: bytes) -> ExtractedDocument:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    doc = ExtractedDocument(page_label="sheet")
    props = wb.properties
    doc.metadata_author = (props.creator or "").strip() if props else ""
    doc.metadata_title = (props.title or "").strip() if props else ""
    doc.metadata_date = _iso(props.created) if props else ""
    for ws in wb.worksheets:
        lines = [f"Sheet: {ws.title}"]
        header: list[str] | None = None
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v).strip() for v in row]
            if not any(values):
                continue
            if header is None:
                header = values
                lines.append(" | ".join(v for v in values if v))
                continue
            pairs = []
            for i, value in enumerate(values):
                if not value:
                    continue
                key = header[i] if i < len(header) and header[i] else ""
                pairs.append(f"{key}: {value}" if key else value)
            lines.append(" | ".join(pairs))
        doc.pages.append("\n".join(lines))
    wb.close()
    if _alpha_count(doc.text) < MIN_ALPHA_TOTAL:
        doc.status = "needs_visual_review"
        doc.status_detail = "The spreadsheet contains little text."
    return doc


def extract_txt(data: bytes) -> ExtractedDocument:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        text = data.decode("utf-8", errors="replace")
    doc = ExtractedDocument(page_label="section", paged=False, pages=[text])
    if "�" in text or _alpha_count(text) < MIN_ALPHA_TOTAL:
        doc.status = "needs_visual_review"
        doc.status_detail = "The text file is empty or has an unreadable encoding."
    return doc


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
    ".txt": extract_txt,
}


def extract(filename: str, data: bytes) -> ExtractedDocument:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        return ExtractedDocument(
            status="needs_visual_review", status_detail=f"Unsupported file type {ext or '(none)'}."
        )
    try:
        return extractor(data)
    except Exception as exc:  # corrupt or encrypted files
        return ExtractedDocument(
            status="needs_visual_review",
            status_detail=f"Text extraction failed: {type(exc).__name__}: {exc}"[:300],
        )


# Near-duplicate detection: a bottom-k sketch of hashed word 5-grams.
SKETCH_SIZE = 128


def _shingle_hashes(text: str, k: int = 5) -> set[int]:
    words = normalize_text(text).split()
    if len(words) < k:
        return {int.from_bytes(hashlib.blake2b(" ".join(words).encode(), digest_size=8).digest(), "big")} if words else set()
    return {
        int.from_bytes(hashlib.blake2b(" ".join(words[i : i + k]).encode(), digest_size=8).digest(), "big")
        for i in range(len(words) - k + 1)
    }


def sketch(text: str) -> list[int]:
    return sorted(_shingle_hashes(text))[:SKETCH_SIZE]


def sketch_similarity(a: list[int], b: list[int]) -> float:
    """Estimate Jaccard similarity from two bottom-k sketches."""
    if not a or not b:
        return 0.0
    sa, sb = set(a), set(b)
    union_bottom = sorted(sa | sb)[:SKETCH_SIZE]
    shared = sum(1 for h in union_bottom if h in sa and h in sb)
    return shared / len(union_bottom)


def jaccard(text_a: str, text_b: str) -> float:
    a, b = _shingle_hashes(text_a), _shingle_hashes(text_b)
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


def containment(text_a: str, text_b: str) -> float:
    """Share of the smaller document's shingles found in the larger one."""
    a, b = _shingle_hashes(text_a), _shingle_hashes(text_b)
    if not a or not b:
        return 0.0
    small, large = (a, b) if len(a) <= len(b) else (b, a)
    return len(small & large) / len(small)
