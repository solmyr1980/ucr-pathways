"""Tests against the known cases from the specification."""

from __future__ import annotations

import io
import json
import re

import pytest

from pec import queries
from pec.classifier import Draft, guardrails
from pec.classifier.rules import classify_passage
from pec.db import Store
from pec.dossier import DossierOptions, dossier_docx, full_report_docx, generate, word_count
from pec.pipeline import analyze, ingest, ingest_folder, rebuild
from pec.samples import CANDIDATE, build


def classify(text):
    draft = classify_passage(text)
    return guardrails.apply(draft) if draft else None


# ---------------------------------------------------------------------------
# Passage classification


def test_title_alone_is_not_l3a():
    d = classify("Head Tutor, 2021-2023")
    assert d.domains[0] == 3 and d.level == "L3A" and d.status == "POTENTIAL"
    d = classify("Interim Head of Social Sciences, 2019-2020")
    assert d.status == "POTENTIAL"


def test_head_tutor_with_institutional_action_is_l3a():
    d = classify("The Head Tutor helps oversee the tutoring system, coordinates training and professional development "
                 "of tutors, participates in evaluation, appraises tutors, and reports findings to the Board of Studies.")
    assert d.level == "L3A" and d.status == "VERIFIED" and 3 in d.domains


@pytest.mark.parametrize("text", [
    "Member of the Board of Studies, 2017-2020",
    "Member, Program Committee, 2015-2017",
    "Extended Executive Board Member, 2020-2022",
])
def test_membership_is_not_leadership(text):
    assert classify(text).level == "L2"


def test_works_council_is_context():
    assert classify("Chair, Works Council, 2016-2018. Vice-Chair, Works Council, 2014-2016.").status == "CONTEXT"


def test_disciplinary_publication_with_student_is_context():
    d = classify("Verhoeven, S. and Jansen, A. (2021). Diane di Prima and the Beat archive. Journal of Beat Studies 9. "
                 "Co-authored with a former student.")
    assert d.status == "CONTEXT" and d.level != "L3B" and d.domains == [1] and d.confidence == "High"


def test_intervention_proposal_is_potential_l3b():
    d = classify("Research question: To what extent can an interdisciplinary educational intervention bridge the disciplines at UCR?")
    assert d.level == "L3B" and d.status == "POTENTIAL" and {5, 6} <= set(d.domains)
    assert any("findings" in m.lower() for m in d.missing_evidence)


def test_teaching_qualification_is_context():
    for text in ("Senior Teaching Qualification (STQ), 2020.", "Senior University Teaching Qualification (SUTQ), 2020."):
        d = classify(text)
        assert d.status == "CONTEXT" and d.domains == [4] and d.level == "Not assigned"


def test_guardrails_lower_llm_overclaims():
    d = guardrails.apply(Draft(passage="Head Tutor, 2021-2023", page_no=1, title="Head Tutor", claim="", domains=[3],
                               level="L3A", level_reason="", evidence_types=["W"], status="VERIFIED", confidence="High"))
    assert d.status == "POTENTIAL"
    d = guardrails.apply(Draft(passage="A monograph on Kerouac and ecology, published 2020.", page_no=1, title="", claim="",
                               domains=[6], level="L3B", level_reason="", evidence_types=["W"], status="VERIFIED", confidence="High"))
    assert d.level == "Not assigned" and d.status == "CONTEXT"


# ---------------------------------------------------------------------------
# Pipeline on the sample corpus


@pytest.fixture(scope="module")
def store(tmp_path_factory):
    folder = build(tmp_path_factory.mktemp("docs"))
    s = Store(tmp_path_factory.mktemp("data") / "test.sqlite3")
    s.set_setting("candidate_name", CANDIDATE)
    ingest_folder(s, folder)
    analyze(s, "rules")
    return s


def by_title(store, fragment):
    return [i for i in queries.evidence(store) if fragment.lower() in i["title"].lower()]


def test_exact_duplicates_become_one_document(store):
    names = [r["filename"] for r in store.q("SELECT filename FROM documents")]
    assert sum("Verhoeven_CV_2024" in n for n in names) == 1
    assert sum(n.endswith(".pdf") and "Handbook" in n for n in names) == 1
    assert queries.counts(store)["duplicate_copies"] == 2


def test_near_duplicate_handbook_needs_review(store):
    kinds = [e["kind"] for e in queries.review_inbox(store)]
    assert "relation" in kinds


def test_unreadable_document_is_kept_for_visual_review(store):
    doc = store.one("SELECT * FROM documents WHERE filename LIKE 'scanned%'")
    assert doc["text_status"] == "needs_visual_review" and not doc["excluded"]
    assert any(e["kind"] == "document" for e in queries.review_inbox(store))


def test_head_tutor_verified_by_cv_and_handbook(store):
    item = by_title(store, "tutoring and advising system")[0]
    assert item["level"] == "L3A" and item["status"] == "VERIFIED" and item["confidence"] == "High"
    assert item["domains"][0] == 3
    files = {s.filename for s in queries.sources_for(store, item["id"])}
    assert any("CV" in f for f in files) and any("Handbook" in f for f in files)


def test_uncorroborated_cv_claim_is_supported_and_flagged(store):
    item = by_title(store, "Led transformation")[0]
    assert item["status"] == "SUPPORTED" and "uncorroborated_claim" in item["review_reasons"]


def test_every_evidence_item_is_traceable(store):
    for item in queries.evidence(store):
        sources = queries.sources_for(store, item["id"])
        assert sources, item["title"]
        for src in sources:
            page = store.one("SELECT text FROM pages WHERE document_id = ? AND page_no = ?", (src.document_id, src.page_no))
            assert re.sub(r"\s+", " ", src.text).strip() in re.sub(r"\s+", " ", page["text"])


def test_requirement_wording_is_conservative(store):
    check = queries.requirement_check(store)
    assert check["statement"].startswith("Current uploaded evidence supports Level 3 claims in")
    assert "qualify" not in json.dumps(check).lower()


def test_decisions_survive_rebuild(store):
    item = by_title(store, "Interim Head")[0] if by_title(store, "Interim Head") else by_title(store, "Social Sciences")[0]
    from pec.pipeline import decide_evidence

    decide_evidence(store, item["fingerprint"], "accept")
    rebuild(store)
    again = queries.evidence(store, "WHERE fingerprint = ?", (item["fingerprint"],))[0]
    assert again["review_state"] == "accepted"


def test_search_finds_evidence_and_documents(store):
    result = queries.search(store, "Head Tutor")
    assert result["evidence"] and result["documents"]


# ---------------------------------------------------------------------------
# Dossier


def test_dossier_default_is_selective_and_conservative(store):
    draft = generate(store, DossierOptions())
    assert draft.overview_header["L3 domains supported"] == "2"
    assert {e["status"] for e in draft.portfolio} == {"VERIFIED"}
    assert 1 <= len(draft.portfolio) <= 15
    titles = " ".join(e["title"] for e in draft.portfolio)
    assert "tutoring and advising system" in titles
    assert not any("di Prima" in e["title"] or "Qualification" in e["title"] for e in draft.portfolio)
    levels = {r["Domain"][:1]: r["Highest supported level"] for r in draft.overview}
    assert "L3A" in levels["3"] and "L3A" in levels["5"]
    assert "L3B" not in levels["5"]
    assert not draft.gaps and not draft.appendix
    for word in ("outstanding", "transformative", "exceptional", "internationally leading", "qualif" + "ies for promotion"):
        assert word not in draft.statement.lower()
    assert 400 <= word_count(draft.statement) <= 1000
    assert "Interdisciplinary" in draft.statement and "potential Level 3B" in draft.statement


def test_dossier_options_and_docx(store):
    import docx

    draft = generate(store, DossierOptions(include_supported=True, include_gaps=True, include_source_list=True, include_appendix=True))
    assert any(e["status"] == "SUPPORTED" for e in draft.portfolio) and draft.gaps and draft.appendix and draft.source_list
    document = docx.Document(io.BytesIO(dossier_docx(draft)))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading 1")]
    assert headings[:3] == ["1. Reflective Statement", "2. Career Framework Overview", "3. Selective Evidence Portfolio"]
    assert "Evidence Still to Locate" in headings
    text = "\n".join(p.text for p in document.paragraphs) + "\n".join(
        c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "confidence" not in text.lower() and "—" not in text


def test_full_report_includes_context(store):
    import docx

    document = docx.Document(io.BytesIO(full_report_docx(store)))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "not the promotion dossier" in text and "Works Council" in text
